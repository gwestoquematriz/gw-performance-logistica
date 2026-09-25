import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference, PieChart, LineChart

print("Loading fleet data...")
with open('c:/Users/renan.alves/.gemini/Projetos/Bancos/frota_dashboard_data.json', encoding='utf-8') as f:
    data = json.load(f)

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Color Palette: Deep Blue / Navy / Corporate Slate
DARK_BLUE = "1B365D"
LIGHT_BLUE = "E8EEF5"
HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1B365D")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="555555")
CARD_TITLE_FONT = Font(name="Calibri", size=9, bold=True, color="555555")
CARD_VAL_FONT = Font(name="Calibri", size=14, bold=True, color="1B365D")
CARD_FILL = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color="000000")
ZEBRA_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='thin', color='DDDDDD')
)

top_double_border = Border(
    top=Side(style='thin', color='000000'),
    bottom=Side(style='double', color='000000')
)

# -------------------------------------------------------------
# TAB 1: RESUMO EXECUTIVO E RANKING DE VEÍCULOS
# -------------------------------------------------------------
ws1 = wb.create_sheet(title="Resumo Geral Frota")
ws1.views.sheetView[0].showGridLines = True

ws1['A1'] = "GW WIRELESS - RELATÓRIO DE GESTÃO DE FROTA E CUSTOS OPERACIONAIS"
ws1['A1'].font = TITLE_FONT
ws1['A2'] = "Consolidação de Combustível (Gasola), Pedágios (Sem Parar), Manutenções e GLPI | Período: 2025 - 2026"
ws1['A2'].font = SUBTITLE_FONT

# KPI Summary Cards (Rows 4-5)
kpis = [
    ("CUSTO TOTAL FROTA", sum(v['custo_total'] for v in data['veiculos']), "R$ #,##0.00", "B4", "C4", "B5", "C5"),
    ("TOTAL COMBUSTÍVEL", sum(v['gasto_combustivel'] for v in data['veiculos']), "R$ #,##0.00", "D4", "E4", "D5", "E5"),
    ("TOTAL MANUTENÇÃO", sum(v['gasto_manutencao'] for v in data['veiculos']), "R$ #,##0.00", "F4", "G4", "F5", "G5"),
    ("TOTAL PEDÁGIO", sum(v['gasto_pedagio'] for v in data['veiculos']), "R$ #,##0.00", "H4", "I4", "H5", "I5"),
    ("KM TOTAL RODADO", sum(v['km_rodado'] for v in data['veiculos']), "#,##0 km", "J4", "K4", "J5", "K5"),
    ("MÉDIA KM/L FROTA", (sum(v['km_rodado'] for v in data['veiculos']) / sum(v['litros'] for v in data['veiculos'])) if sum(v['litros'] for v in data['veiculos']) > 0 else 0, "0.00 km/L", "L4", "M4", "L5", "M5"),
]

for title, val, num_fmt, t_start, t_end, v_start, v_end in kpis:
    ws1.merge_cells(f"{t_start}:{t_end}")
    ws1.merge_cells(f"{v_start}:{v_end}")
    cell_t = ws1[t_start]
    cell_t.value = title
    cell_t.font = CARD_TITLE_FONT
    cell_t.alignment = Alignment(horizontal="center", vertical="center")
    cell_t.fill = CARD_FILL
    
    cell_v = ws1[v_start]
    cell_v.value = val
    cell_v.font = CARD_VAL_FONT
    cell_v.alignment = Alignment(horizontal="center", vertical="center")
    cell_v.number_format = num_fmt
    cell_v.fill = CARD_FILL

# Table Headers
headers_tab1 = [
    "Placa", "Modelo", "Apelido", "Status", "Filial Operação", "Motorista Atual",
    "Gasto Combustível (R$)", "Litros Consumidos", "Km Rodado Total", "Consumo Médio (Km/L)",
    "Gasto Pedágio (R$)", "Qtd Pedágios", "Gasto Manutenção (R$)", "Custo Total (R$)", "Custo Médio (R$/Km)"
]

start_row = 8
for col_idx, h in enumerate(headers_tab1, 1):
    cell = ws1.cell(row=start_row, column=col_idx, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

row_curr = start_row + 1
for v in data['veiculos']:
    c_km = (v['custo_total'] / v['km_rodado']) if v['km_rodado'] > 0 else 0
    row_data = [
        v['placa'],
        v['modelo'],
        v['apelido'],
        v['status'],
        v['filial'],
        v['motorista'],
        v['gasto_combustivel'],
        v['litros'],
        v['km_rodado'],
        v['media_kml'],
        v['gasto_pedagio'],
        v['qtd_pedagio'],
        v['gasto_manutencao'],
        v['custo_total'],
        c_km
    ]
    for col_idx, val in enumerate(row_data, 1):
        cell = ws1.cell(row=row_curr, column=col_idx, value=val)
        cell.border = thin_border
        cell.font = Font(name="Calibri", size=10)
        
        # Formats
        if col_idx in [1, 3, 4, 5]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [2, 6]:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        elif col_idx in [7, 11, 13, 14]:
            cell.number_format = "R$ #,##0.00"
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx in [8, 9]:
            cell.number_format = "#,##0.0"
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx == 10:
            cell.number_format = "0.00"
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx == 12:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx == 15:
            cell.number_format = "R$ 0.00"
            cell.alignment = Alignment(horizontal="right", vertical="center")
            
        if (row_curr - start_row) % 2 == 0:
            cell.fill = ZEBRA_FILL
            
    row_curr += 1

# Total Row
total_row = row_curr
ws1.cell(row=total_row, column=1, value="TOTAL GERAL").font = TOTAL_FONT
ws1.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
ws1.cell(row=total_row, column=1).alignment = Alignment(horizontal="center", vertical="center")

# Formulas
ws1.cell(row=total_row, column=7, value=f"=SUM(G9:G{total_row-1})").number_format = "R$ #,##0.00"
ws1.cell(row=total_row, column=8, value=f"=SUM(H9:H{total_row-1})").number_format = "#,##0.0"
ws1.cell(row=total_row, column=9, value=f"=SUM(I9:I{total_row-1})").number_format = "#,##0.0"
ws1.cell(row=total_row, column=10, value=f"=I{total_row}/H{total_row}").number_format = "0.00"
ws1.cell(row=total_row, column=11, value=f"=SUM(K9:K{total_row-1})").number_format = "R$ #,##0.00"
ws1.cell(row=total_row, column=12, value=f"=SUM(L9:L{total_row-1})").number_format = "#,##0"
ws1.cell(row=total_row, column=13, value=f"=SUM(M9:M{total_row-1})").number_format = "R$ #,##0.00"
ws1.cell(row=total_row, column=14, value=f"=SUM(N9:N{total_row-1})").number_format = "R$ #,##0.00"
ws1.cell(row=total_row, column=15, value=f"=N{total_row}/I{total_row}").number_format = "R$ 0.00"

for col_idx in range(1, 16):
    c = ws1.cell(row=total_row, column=col_idx)
    c.font = TOTAL_FONT
    c.fill = TOTAL_FILL
    c.border = top_double_border
    if col_idx >= 7:
        c.alignment = Alignment(horizontal="right", vertical="center")

# -------------------------------------------------------------
# TAB 2: EVOLUÇÃO MENSAL DE CUSTOS (KARDEX)
# -------------------------------------------------------------
ws2 = wb.create_sheet(title="Evolução Mensal Custos")
ws2.views.sheetView[0].showGridLines = True

ws2['A1'] = "HISTÓRICO MENSAL CONSOLIDADO DE CUSTOS POR VEÍCULO"
ws2['A1'].font = TITLE_FONT
ws2['A2'] = "Despesas de Abastecimento, Pedágio, Peças e Serviços de Manutenção por Mês"
ws2['A2'].font = SUBTITLE_FONT

headers_tab2 = ["Ano-Mês", "Filial", "Placa", "Apelido", "Categoria de Despesa", "Valor Total (R$)"]
for col_idx, h in enumerate(headers_tab2, 1):
    cell = ws2.cell(row=4, column=col_idx, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")

r_idx = 5
for row in data['custos_mensais']:
    ws2.cell(row=r_idx, column=1, value=row['ano_mes']).alignment = Alignment(horizontal="center")
    ws2.cell(row=r_idx, column=2, value=row['filial']).alignment = Alignment(horizontal="center")
    ws2.cell(row=r_idx, column=3, value=row['placa']).alignment = Alignment(horizontal="center")
    ws2.cell(row=r_idx, column=4, value=row['apelido']).alignment = Alignment(horizontal="left")
    ws2.cell(row=r_idx, column=5, value=row['categoria']).alignment = Alignment(horizontal="center")
    c_val = ws2.cell(row=r_idx, column=6, value=row['valor'])
    c_val.number_format = "R$ #,##0.00"
    c_val.alignment = Alignment(horizontal="right")
    
    for c_i in range(1, 7):
        ws2.cell(row=r_idx, column=c_i).border = thin_border
        ws2.cell(row=r_idx, column=c_i).font = Font(name="Calibri", size=10)
    if (r_idx - 5) % 2 == 1:
        for c_i in range(1, 7):
            ws2.cell(row=r_idx, column=c_i).fill = ZEBRA_FILL
    r_idx += 1

# -------------------------------------------------------------
# TAB 3: HISTÓRICO DE MANUTENÇÕES E OFICINAS
# -------------------------------------------------------------
ws3 = wb.create_sheet(title="Manutenções & Oficinas")
ws3.views.sheetView[0].showGridLines = True

ws3['A1'] = "DETALHAMENTO DE ORDENS DE MANUTENÇÃO, PEÇAS E SERVIÇOS"
ws3['A1'].font = TITLE_FONT
ws3['A2'] = "Registros analíticos de manutenções preventivas e corretivas da frota"
ws3['A2'].font = SUBTITLE_FONT

headers_tab3 = ["ID", "Data", "Mês Ref.", "Placa", "Veículo", "Grupo", "Filial", "Cidade", "Valor (R$)", "Descrição do Reparo / Serviço"]
for col_idx, h in enumerate(headers_tab3, 1):
    cell = ws3.cell(row=4, column=col_idx, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")

r_idx = 5
for row in data['manutencoes']:
    ws3.cell(row=r_idx, column=1, value=row['id']).alignment = Alignment(horizontal="center")
    ws3.cell(row=r_idx, column=2, value=row['data']).alignment = Alignment(horizontal="center")
    ws3.cell(row=r_idx, column=3, value=row['ano_mes']).alignment = Alignment(horizontal="center")
    ws3.cell(row=r_idx, column=4, value=row['placa']).alignment = Alignment(horizontal="center")
    ws3.cell(row=r_idx, column=5, value=row['veiculo']).alignment = Alignment(horizontal="left")
    ws3.cell(row=r_idx, column=6, value=row['grupo']).alignment = Alignment(horizontal="center")
    ws3.cell(row=r_idx, column=7, value=row['filial']).alignment = Alignment(horizontal="center")
    ws3.cell(row=r_idx, column=8, value=row['cidade']).alignment = Alignment(horizontal="left")
    c_val = ws3.cell(row=r_idx, column=9, value=row['valor'])
    c_val.number_format = "R$ #,##0.00"
    c_val.alignment = Alignment(horizontal="right")
    ws3.cell(row=r_idx, column=10, value=row['descricao']).alignment = Alignment(horizontal="left")
    
    for c_i in range(1, 11):
        ws3.cell(row=r_idx, column=c_i).border = thin_border
        ws3.cell(row=r_idx, column=c_i).font = Font(name="Calibri", size=9)
    if (r_idx - 5) % 2 == 1:
        for c_i in range(1, 11):
            ws3.cell(row=r_idx, column=c_i).fill = ZEBRA_FILL
    r_idx += 1

# -------------------------------------------------------------
# TAB 4: COMBUSTÍVEL GASOLA MENSAL
# -------------------------------------------------------------
ws4 = wb.create_sheet(title="Combustível Gasola")
ws4.views.sheetView[0].showGridLines = True

ws4['A1'] = "EFICIÊNCIA DE COMBUSTÍVEL E CONSUMO (PLATAFORMA GASOLA)"
ws4['A1'].font = TITLE_FONT
ws4['A2'] = "Consumo de diesel, gasolina e arla, km rodado e eficiência km/litro"
ws4['A2'].font = SUBTITLE_FONT

headers_tab4 = ["Ano-Mês", "Placa", "Qtd Abastecimentos", "Valor Pago (R$)", "Litros Consumidos", "Km Rodado", "Preço Médio Litro (R$)", "Média Km/L"]
for col_idx, h in enumerate(headers_tab4, 1):
    cell = ws4.cell(row=4, column=col_idx, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")

r_idx = 5
for row in data['combustivel_mensal']:
    ws4.cell(row=r_idx, column=1, value=row['ano_mes']).alignment = Alignment(horizontal="center")
    ws4.cell(row=r_idx, column=2, value=row['placa']).alignment = Alignment(horizontal="center")
    ws4.cell(row=r_idx, column=3, value=row['qtd']).alignment = Alignment(horizontal="right")
    c_val = ws4.cell(row=r_idx, column=4, value=row['valor'])
    c_val.number_format = "R$ #,##0.00"
    c_val.alignment = Alignment(horizontal="right")
    c_lit = ws4.cell(row=r_idx, column=5, value=row['litros'])
    c_lit.number_format = "#,##0.0"
    c_lit.alignment = Alignment(horizontal="right")
    c_km = ws4.cell(row=r_idx, column=6, value=row['km'])
    c_km.number_format = "#,##0.0"
    c_km.alignment = Alignment(horizontal="right")
    c_pr = ws4.cell(row=r_idx, column=7, value=row['preco_medio'])
    c_pr.number_format = "R$ 0.000"
    c_pr.alignment = Alignment(horizontal="right")
    c_kml = ws4.cell(row=r_idx, column=8, value=row['kml'])
    c_kml.number_format = "0.00"
    c_kml.alignment = Alignment(horizontal="right")
    
    for c_i in range(1, 9):
        ws4.cell(row=r_idx, column=c_i).border = thin_border
        ws4.cell(row=r_idx, column=c_i).font = Font(name="Calibri", size=10)
    if (r_idx - 5) % 2 == 1:
        for c_i in range(1, 9):
            ws4.cell(row=r_idx, column=c_i).fill = ZEBRA_FILL
    r_idx += 1

# -------------------------------------------------------------
# TAB 5: TRANSFERÊNCIAS GLPI
# -------------------------------------------------------------
ws5 = wb.create_sheet(title="Transferências GLPI")
ws5.views.sheetView[0].showGridLines = True

ws5['A1'] = "CHAMADOS GLPI - TRANSFERÊNCIA DE DESPESAS DE ABASTECIMENTO"
ws5['A1'].font = TITLE_FONT
ws5['A2'] = "Rateio e transferência de abastecimentos entre Matriz e Filiais"
ws5['A2'].font = SUBTITLE_FONT

headers_tab5 = ["ID Chamado", "Título do Chamado", "Status", "Data Abertura", "Solicitante", "Empresa Origem", "Empresa Destino", "Motivo"]
for col_idx, h in enumerate(headers_tab5, 1):
    cell = ws5.cell(row=4, column=col_idx, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")

r_idx = 5
for row in data['glpi_transfers']:
    ws5.cell(row=r_idx, column=1, value=row['id']).alignment = Alignment(horizontal="center")
    ws5.cell(row=r_idx, column=2, value=row['titulo']).alignment = Alignment(horizontal="left")
    ws5.cell(row=r_idx, column=3, value=row['status']).alignment = Alignment(horizontal="center")
    ws5.cell(row=r_idx, column=4, value=row['data']).alignment = Alignment(horizontal="center")
    ws5.cell(row=r_idx, column=5, value=row['solicitante']).alignment = Alignment(horizontal="left")
    ws5.cell(row=r_idx, column=6, value=row['origem']).alignment = Alignment(horizontal="left")
    ws5.cell(row=r_idx, column=7, value=row['destino']).alignment = Alignment(horizontal="left")
    ws5.cell(row=r_idx, column=8, value=row['motivo']).alignment = Alignment(horizontal="left")
    
    for c_i in range(1, 9):
        ws5.cell(row=r_idx, column=c_i).border = thin_border
        ws5.cell(row=r_idx, column=c_i).font = Font(name="Calibri", size=9)
    if (r_idx - 5) % 2 == 1:
        for c_i in range(1, 9):
            ws5.cell(row=r_idx, column=c_i).fill = ZEBRA_FILL
    r_idx += 1

# Auto-adjust column widths on all sheets
for ws in wb.worksheets:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # ignore merged banner title
            if cell.row in [1, 2, 4, 5] and ws.title == "Resumo Geral Frota":
                continue
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

# Add Chart in Tab 1
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Top 10 Veículos por Custo Total (R$)"
chart.y_axis.title = "Custo Total (R$)"
chart.x_axis.title = "Veículo (Placa)"
chart.height = 14
chart.width = 24

data_ref = Reference(ws1, min_col=14, min_row=8, max_row=18)
cats_ref = Reference(ws1, min_col=1, min_row=9, max_row=18)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
chart.legend = None

ws1.add_chart(chart, "B36")

excel_path = 'c:/Users/renan.alves/.gemini/Projetos/Bancos/Relatorio_Gestao_Frota_Custos.xlsx'
wb.save(excel_path)
print(f"Excel report generated successfully at: {excel_path}")
